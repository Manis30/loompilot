import os
from pathlib import Path
from typing import Any, Dict, List, Optional
import pandas as pd

from app.core.config import settings
from app.core.logging import get_logger

logger = get_logger(__name__)


class ExcelDataLoader:
    """
    Safe, read-only loader for LoomPilot Excel datasets.
    Guarantees that original datasets are never modified or overwritten.
    """

    def __init__(self, dataset_dir: Optional[Path] = None):
        self.dataset_dir = Path(dataset_dir) if dataset_dir else settings.dataset_path
        if not self.dataset_dir.exists():
            logger.warning(f"Dataset directory does not exist: {self.dataset_dir}")

    def list_available_files(self) -> List[str]:
        """List all Excel files present in the dataset directory."""
        if not self.dataset_dir.exists():
            return []
        return sorted([
            f for f in os.listdir(self.dataset_dir)
            if f.endswith('.xlsx') or f.endswith('.xls')
        ])

    def load_sheet(self, filename: str, sheet_name: Optional[str] = None) -> pd.DataFrame:
        """
        Load a specific sheet from an Excel file into a Pandas DataFrame.
        Read-only, preserves original file.
        """
        file_path = self.dataset_dir / filename
        if not file_path.exists():
            raise FileNotFoundError(f"Excel file not found at: {file_path}")

        try:
            # openpyxl engine used for modern .xlsx files
            if sheet_name is not None:
                df = pd.read_excel(file_path, sheet_name=sheet_name, engine="openpyxl")
            else:
                df = pd.read_excel(file_path, engine="openpyxl")
            return df
        except Exception as e:
            logger.error(f"Error loading Excel file {file_path}: {e}")
            raise RuntimeError(f"Failed to read Excel file {filename}: {e}") from e

    def get_file_sheets(self, filename: str) -> List[str]:
        """Return list of sheet names within an Excel file."""
        file_path = self.dataset_dir / filename
        if not file_path.exists():
            raise FileNotFoundError(f"Excel file not found at: {file_path}")

        excel = pd.ExcelFile(file_path, engine="openpyxl")
        return excel.sheet_names

    def load_market_prices(self) -> pd.DataFrame:
        """Load primary market prices dataset (LoomPilot_MarketPrices.xlsx)."""
        return self.load_sheet("LoomPilot_MarketPrices.xlsx", sheet_name="MarketPrices")

    def load_materials(self) -> pd.DataFrame:
        """Load raw materials master dataset (LoomPilot_Materials.xlsx)."""
        return self.load_sheet("LoomPilot_Materials.xlsx", sheet_name="Materials")

    def load_suppliers(self) -> pd.DataFrame:
        """Load suppliers directory dataset (LoomPilot_Suppliers.xlsx)."""
        return self.load_sheet("LoomPilot_Suppliers.xlsx", sheet_name="Suppliers")

    def load_order_materials(self) -> pd.DataFrame:
        """Load order material locked rates dataset (LoomPilot_OrderMaterials.xlsx)."""
        return self.load_sheet("LoomPilot_OrderMaterials.xlsx", sheet_name="OrderMaterials")

    def load_orders(self) -> pd.DataFrame:
        """Load orders dataset (LoomPilot_Orders.xlsx)."""
        return self.load_sheet("LoomPilot_Orders.xlsx", sheet_name="Orders")

    def load_quotations(self) -> pd.DataFrame:
        """Load quotations dataset (LoomPilot_Quotations.xlsx)."""
        return self.load_sheet("LoomPilot_Quotations.xlsx", sheet_name="Quotations")

    def load_products(self) -> pd.DataFrame:
        """Load products master dataset (LoomPilot_Products.xlsx)."""
        return self.load_sheet("LoomPilot_Products.xlsx", sheet_name="Products")

    def load_machines(self) -> pd.DataFrame:
        """Load machines master dataset (LoomPilot_Machines.xlsx)."""
        return self.load_sheet("LoomPilot_Machines.xlsx", sheet_name="Machines")

    def load_machine_status(self) -> pd.DataFrame:
        """Load machine status history dataset (LoomPilot_MachineStatus.xlsx)."""
        return self.load_sheet("LoomPilot_MachineStatus.xlsx", sheet_name="MachineStatus")

    def load_production_reports(self) -> pd.DataFrame:
        """Load daily production reports dataset (LoomPilot_ProductionReports.xlsx)."""
        return self.load_sheet("LoomPilot_ProductionReports.xlsx", sheet_name="ProductionReports")

    def load_attendance(self) -> pd.DataFrame:
        """Load employee attendance records dataset (LoomPilot_Attendance.xlsx)."""
        return self.load_sheet("LoomPilot_Attendance.xlsx", sheet_name="Attendance")

    def load_employees(self) -> pd.DataFrame:
        """Load employees directory dataset (LoomPilot_Employees.xlsx)."""
        return self.load_sheet("LoomPilot_Employees.xlsx", sheet_name="Employees")

    def load_supervisors(self) -> pd.DataFrame:
        """Load supervisors directory dataset (LoomPilot_Supervisors.xlsx)."""
        return self.load_sheet("LoomPilot_Supervisors.xlsx", sheet_name="Supervisors")

    def load_clients(self) -> pd.DataFrame:
        """Load clients directory dataset (LoomPilot_Clients.xlsx)."""
        return self.load_sheet("LoomPilot_Clients.xlsx", sheet_name="Clients")

    def get_dataset_inspection_summary(self) -> Dict[str, Any]:
        """
        Inspect all available Excel files in the dataset folder,
        returning sheet names, dimensions, column names, and row counts.
        Caches results for high performance.
        """
        if hasattr(self, "_cached_inspection") and self._cached_inspection:
            return self._cached_inspection

        import openpyxl

        summary: Dict[str, Any] = {}
        for fname in self.list_available_files():
            file_path = self.dataset_dir / fname
            try:
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                sheets_data = []
                for s_name in wb.sheetnames:
                    ws = wb[s_name]
                    headers = []
                    row_count = 0
                    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                        if row_idx == 0:
                            headers = [str(c) for c in row if c is not None]
                        if any(c is not None for c in row):
                            row_count += 1
                    sheets_data.append({
                        "sheetName": s_name,
                        "rowCount": max(0, row_count - 1),
                        "columnCount": len(headers),
                        "columns": headers,
                    })
                wb.close()
                summary[fname] = sheets_data
            except Exception as e:
                logger.warning(f"Could not inspect {fname}: {e}")
                summary[fname] = [{"error": str(e)}]

        self._cached_inspection = summary
        return summary
